"""Explicit real model / PostgreSQL / OSS checks using disposable fictitious data."""

import argparse
import asyncio
import json
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from uuid import uuid4

import psycopg
from psycopg import sql
from alembic import command
from alembic.config import Config
from sqlalchemy import create_engine, select, func
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from docx import Document as Docx
from fpdf import FPDF
import httpx
from fastapi import FastAPI

from common.account.models import User
from common.settings import POSTGRES_DATABASE
from common.storage.adapter import OssObjectStorage
from common.config.config import dashscope_model_name, dashscope_rerank_model
from common.integrations.retrieval_models import (
    make_dashscope_embeddings,
    rerank_indices,
)
from products.contract.models import ContractKnowledgeDocumentModel as Document
from products.contract.models import ContractKnowledgeChunkModel as Chunk
from products.contract.knowledge import ingest_document
from products.contract.retrieval import retrieve_references, validate_vectors
from scripts.contract.test_contract_migration_postgres import (
    BACKEND_ROOT,
    _url,
    _psycopg_url,
    _drop_database,
)


def model_check():
    texts = ["企业采购规范：付款期限原则上不超过45天。", "供应商交付后七日内安排验收。"]
    model = make_dashscope_embeddings(dimensions=1024, timeout=25, max_retries=0)
    validate_vectors(model.embed_documents(texts), 2)
    validate_vectors([model.embed_query("采购合同付款期限90天是否超过企业规范")], 1)
    indices = rerank_indices("采购合同付款期限", texts, top_n=2, timeout=25)
    print(
        json.dumps(
            {
                "real_embedding": dashscope_model_name(),
                "dimensions": 1024,
                "real_rerank": dashscope_rerank_model(),
                "reranked_indices": indices,
                "postgresql_oss_chain": "NOT RUN in --models-only mode",
            },
            ensure_ascii=False,
        ),
        flush=True,
    )


async def verify(database_name, *, review=False):
    engine = create_async_engine(
        _url(database_name).set(drivername="postgresql+asyncpg")
    )
    factory = async_sessionmaker(engine, expire_on_commit=False)
    storage, owned_keys = OssObjectStorage(), []
    try:
        async with factory() as session, session.begin():
            session.add_all(
                [
                    User(
                        user_id=i,
                        username=f"kb-test-{i}",
                        email=f"kb{i}@test.invalid",
                        password="fictitious",
                    )
                    for i in (1, 2)
                ]
            )
        document_ids = []
        for owner_id, kind in ((1, "docx"), (2, "docx"), (1, "pdf")):
            document_id = uuid4()
            key = f"contract-knowledge/test/{database_name}/{document_id}/v1/original.{kind}"
            if kind == "docx":
                stream = BytesIO()
                document = Docx()
                document.add_paragraph(
                    "企业付款不超过45天。"
                    if owner_id == 1
                    else "用户B私有规范，不得向其他用户召回。"
                )
                document.save(stream)
                data = stream.getvalue()
            else:
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Helvetica", size=12)
                pdf.cell(
                    text="Enterprise procurement: payment term should not exceed 45 days."
                )
                data = bytes(pdf.output())
            owned_keys.append(key)
            await asyncio.to_thread(
                storage.put,
                key,
                data,
                content_type=(
                    "application/pdf"
                    if kind == "pdf"
                    else "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
            )
            async with factory() as session, session.begin():
                session.add(
                    Document(
                        id=document_id,
                        user_id=owner_id,
                        filename=f"企业规范-{owner_id}.{kind}",
                        source_key=key,
                        content_hash=uuid4().hex * 2,
                        byte_count=len(data),
                        status="uploaded",
                    )
                )
            await asyncio.gather(
                ingest_document(document_id, 1, factory=factory, storage=storage),
                ingest_document(document_id, 1, factory=factory, storage=storage),
            )
            async with factory() as session:
                stored = await session.get(Document, document_id)
                assert stored.status == "ready", stored.error
                assert stored.attempt == 1
                assert stored.chunk_count == await session.scalar(
                    select(func.count())
                    .select_from(Chunk)
                    .where(Chunk.document_id == document_id)
                )
            document_ids.append(document_id)
        references = await retrieve_references(factory, 1, "企业采购付款期限90天")
        own_ids = {str(document_ids[0]), str(document_ids[2])}
        assert references and all(
            r["owner_id"] == 1 and r["document_id"] in own_ids for r in references
        )
        pdf_references = await retrieve_references(
            factory, 1, "Enterprise procurement payment term 45 days"
        )
        assert any(
            r["document_id"] == str(document_ids[2]) and r["locator"]["pages"] == [1]
            for r in pdf_references
        )
        async with factory() as session, session.begin():
            doc = await session.get(Document, document_ids[0])
            doc.status, doc.generation = "uploaded", 2
        await asyncio.gather(
            ingest_document(document_ids[0], 2, factory=factory, storage=storage),
            ingest_document(document_ids[0], 2, factory=factory, storage=storage),
        )
        updated = await retrieve_references(factory, 1, "企业采购付款期限90天")
        first = next(r for r in updated if r["document_id"] == str(document_ids[0]))
        assert first["index_generation"] == 2
        assert first["chunk_id"] not in {r["chunk_id"] for r in references}
        async with factory() as session:
            stored = await session.get(Document, document_ids[0])
            assert stored.attempt == 2 and stored.chunk_count == await session.scalar(
                select(func.count())
                .select_from(Chunk)
                .where(Chunk.document_id == stored.id)
            )
        print(
            "REAL OSS -> PDF/DOCX Parser -> Embedding -> pgvector + BM25 -> RRF -> Rerank; owner isolation, concurrent claims, idempotence and reindex: PASSED",
            flush=True,
        )
        if review:
            await verify_review_chain(factory, owned_keys)
    finally:
        # The review mode creates originals through real HTTP APIs in this disposable DB.
        from products.contract.models import ContractVersionModel

        async with factory() as session:
            owned_keys.extend(
                (await session.scalars(select(ContractVersionModel.source_key))).all()
            )
            owned_keys.extend(
                (await session.scalars(select(Document.source_key))).all()
            )
        await engine.dispose()
        for key in dict.fromkeys(owned_keys):
            await asyncio.to_thread(storage.delete, key)


async def verify_review_chain(factory, owned_keys):
    """Real models throughout; only authentication and DB factories target fixtures."""
    from common.auth.dependencies import require_user
    from common.database import get_db
    from products.contract import api, db, knowledge
    from products.contract.document import contract_jobs
    from products.contract.agents.review import prompts
    from products.contract.models import ContractVersionModel
    from uuid import UUID
    from time import perf_counter

    app, caller = FastAPI(), [1]
    app.include_router(api.router)
    app.dependency_overrides[require_user] = lambda: SimpleNamespace(user_id=caller[0])

    async def sessions():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_db] = sessions

    def docx(texts):
        doc, stream = Docx(), BytesIO()
        for text in texts:
            doc.add_paragraph(text)
        doc.save(stream)
        return stream.getvalue()

    async def data(response):
        assert response.status_code == 200, f"HTTP {response.status_code}"
        return response.json()["data"]

    with patch.object(contract_jobs, "SessionLocal", factory), patch.object(
        knowledge, "SessionLocal", factory
    ):
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app), base_url="http://test", timeout=300
        ) as client:
            policy = await data(
                await client.post(
                    "/api/contract-policies",
                    json=dict(
                        name="虚构采购保障规则",
                        requirement="买方不得在没有交付和退款保障的情况下全额预付款。",
                        severity="high",
                        applicable_party_role="buyer",
                        applicable_contract_types=["procurement"],
                    ),
                )
            )
            uploaded = await data(
                await client.post(
                    "/api/contract-knowledge/documents/upload",
                    data={"category": "legal_policy"},
                    files={
                        "file": (
                            "虚构企业付款规范.docx",
                            docx(
                                [
                                    "企业采购付款参考：全额付款之前应取得明确交付和退款保障。",
                                    "历史案例和企业参考规范不是法律结论，需法务确认。",
                                ]
                            ),
                        )
                    },
                )
            )
            did = uploaded["document"]["id"]
            listed = await data(await client.get("/api/contract-knowledge/documents"))
            indexed = next(d for d in listed["documents"] if d["id"] == did)
            assert indexed["status"] == "ready", indexed.get("error")
            customer = await data(
                await client.post(
                    "/api/customers",
                    json={"name": "虚构采购客户", "email": "p4@test.invalid"},
                )
            )
            contract = await data(
                await client.post(
                    "/api/contracts",
                    json={"title": "P4 全链路虚构合同", "customer_id": customer["id"]},
                )
            )
            cid = contract["id"]
            scope = dict(
                our_party="甲方", party_role="buyer", contract_type="procurement"
            )
            upload_started = perf_counter()
            # Test-only decision guidance: REAL model with auto tool choice;
            # no mocked messages, model results, tool executor or retrieval.
            with patch.object(
                prompts,
                "DECISION_INSTRUCTIONS",
                prompts.DECISION_INSTRUCTIONS
                + "\n本次验收资料已经充分，请直接开始审查，无需额外工具。",
            ):
                first = await data(
                    await client.post(
                        f"/api/contracts/{cid}/versions/upload",
                        data=scope,
                        files={
                            "file": (
                                "虚构合同.docx",
                                docx(
                                    [
                                        "采购合同",
                                        "甲方：虚构采购方。乙方：虚构供应商。",
                                        "甲方应在签署后10日内全额支付100000元。",
                                        "乙方应在收到款项后90日内交付。",
                                        "任何情况下均不支持退款。",
                                        "甲方赔偿责任不设上限。",
                                    ]
                                ),
                            )
                        },
                    )
                )
            vid, rid = first["version_id"], first["run_id"]
            base = f"/api/contracts/{cid}/versions/{vid}"
            report = await data(await client.get(base + "/analysis"))
            assert report["status"] == "completed", report.get("error")
            assert report["context_snapshot"]["quality_gate"]["status"] == "passed"
            assert report["context_snapshot"]["policies"][0]["version"] == 1
            assert report["context_snapshot"]["references"] and report["risks"]
            assert any(
                r["policy_ids"] for r in report["risks"]
            ), "Live Review did not use the applicable enterprise policy"
            assert any(
                r["reference_citations"] for r in report["risks"]
            ), "Live Review did not cite relevant retrieved source text"
            actual_models = [
                call["model"]
                for call in report["context_snapshot"]["review_trace"]["model_calls"]
                if call.get("model")
            ]
            assert actual_models, "No actual model identity recorded"
            assert not report["context_snapshot"]["review_trace"][
                "tool_calls"
            ], "No-tool live case called a tool"
            assert (
                report["context_snapshot"]["review_trace"]["graph_version"]
                == "review-graph-v1"
            )
            async with factory() as session:
                first_facts = await db.get_document_facts(session, UUID(vid))
            print(
                json.dumps(
                    {
                        "real_review_models": list(dict.fromkeys(actual_models)),
                        "quality_gate": "passed",
                        "risk_count": len(report["risks"]),
                        "fixture": "DOCX / no-tool (test-only decision guidance)",
                        "graph_seconds": report["context_snapshot"]["review_trace"][
                            "elapsed_seconds"
                        ],
                        "token_usage": [
                            c.get("usage")
                            for c in report["context_snapshot"]["review_trace"][
                                "model_calls"
                            ]
                        ],
                        "upload_review_seconds": round(
                            perf_counter() - upload_started, 3
                        ),
                        "coverage": report["facts"]["coverage"]["ratio"],
                        "model_calls": len(
                            report["context_snapshot"]["review_trace"]["model_calls"]
                        ),
                        "tool_calls": len(
                            report["context_snapshot"]["review_trace"].get(
                                "tool_calls", []
                            )
                        ),
                        "quality_input_count": report["context_snapshot"][
                            "quality_gate"
                        ]["input_count"],
                        "quality_verified_count": report["context_snapshot"][
                            "quality_gate"
                        ]["verified_count"],
                    },
                    ensure_ascii=False,
                ),
                flush=True,
            )
            await data(
                await client.patch(
                    f"/api/contract-policies/{policy['id']}",
                    json={
                        "expected_version": 1,
                        "requirement": "买方全额付款前应有验收保障，不得接受任何情况下不退款。",
                    },
                )
            )
            await data(
                await client.post(f"/api/contract-knowledge/documents/{did}/reindex")
            )
            rerun_started = perf_counter()
            with patch.object(
                prompts,
                "DECISION_INSTRUCTIONS",
                prompts.DECISION_INSTRUCTIONS
                + "\n本次验收需额外检索企业全额预付款的交付和退款保障，请先调用一次 search_knowledge，"
                "收到结果后如果足够就开始审查，不重复调用。查询内容由你根据当前合同决定。",
            ):
                await data(await client.post(base + "/analyze", json=scope))
            updated = await data(await client.get(base + "/analysis"))
            assert updated["status"] == "completed", updated.get("error")
            assert updated["context_snapshot"]["policies"][0]["version"] == 2
            assert any(
                r["document_id"] == did and r["index_generation"] == 2
                for r in updated["context_snapshot"]["references"]
            )
            trace = updated["context_snapshot"]["review_trace"]
            searches = [
                c for c in trace["tool_calls"] if c["name"] == "search_knowledge"
            ]
            assert searches and all(c["status"] == "passed" for c in searches)
            assert any(
                c["reference_count"] for c in searches
            ), "Real tool did not retrieve sources"
            assert sum(c["phase"] == "decide" for c in trace["model_calls"]) >= 2
            assert any(
                e["kind"] == "supplementary"
                for e in updated["context_snapshot"]["retrieval_events"]
            )
            assert any(r["reference_citations"] for r in updated["risks"])
            print(
                json.dumps(
                    dict(
                        fixture="DOCX / real-model tool (test-only decision guidance)",
                        coverage=updated["facts"]["coverage"]["ratio"],
                        risk_count=len(updated["risks"]),
                        quality_input_count=updated["context_snapshot"]["quality_gate"][
                            "input_count"
                        ],
                        quality_verified_count=updated["context_snapshot"][
                            "quality_gate"
                        ]["verified_count"],
                        model_calls=len(trace["model_calls"]),
                        tool_calls=len(trace["tool_calls"]),
                        tools=[
                            dict(
                                name=c["name"],
                                status=c["status"],
                                reference_count=c.get("reference_count"),
                            )
                            for c in trace["tool_calls"]
                        ],
                        graph_seconds=trace["elapsed_seconds"],
                        rerun_seconds=round(perf_counter() - rerun_started, 3),
                        token_usage=[c.get("usage") for c in trace["model_calls"]],
                    ),
                    ensure_ascii=False,
                ),
                flush=True,
            )
            async with factory() as session:
                assert await db.get_document_facts(session, UUID(vid)) == first_facts
            await data(await client.delete(f"/api/contract-knowledge/documents/{did}"))
            selected = await data(
                await client.post(base + f"/analysis-runs/{rid}/select")
            )
            assert (
                selected["context_snapshot"] == report["context_snapshot"]
                and selected["result"] == report["result"]
            )
            citation = next(
                r
                for r in report["context_snapshot"]["references"]
                if r["document_id"] == did
            )
            historical = await data(
                await client.get(
                    base + f"/analysis-runs/{rid}/references/{citation['chunk_id']}"
                )
            )
            assert historical == citation
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Helvetica", size=12)
            pdf.multi_cell(
                0,
                8,
                text="Procurement agreement. Buyer pays the full price within 10 days after signing.\nSupplier delivers 90 days after payment. No refunds under any circumstances.",
            )
            uploaded_pdf = await data(
                await client.post(
                    f"/api/contracts/{cid}/versions/upload",
                    data=dict(
                        our_party="Buyer",
                        party_role="buyer",
                        contract_type="procurement",
                    ),
                    files={
                        "file": (
                            "fictitious.pdf",
                            bytes(pdf.output()),
                            "application/pdf",
                        )
                    },
                )
            )
            pdf_report = await data(
                await client.get(
                    f"/api/contracts/{cid}/versions/{uploaded_pdf['version_id']}/analysis"
                )
            )
            assert pdf_report["status"] == "completed", pdf_report.get("error")
            assert pdf_report["context_snapshot"]["quality_gate"]["status"] == "passed"
            assert pdf_report["canonical_clauses"] and all(
                c["pages"] == [1] for c in pdf_report["canonical_clauses"]
            )
            caller[0] = 2
            assert (await client.get(base + "/analysis")).status_code == 404
            assert (
                await client.get(
                    base + f"/analysis-runs/{rid}/references/{citation['chunk_id']}"
                )
            ).status_code == 404
            print(
                "REAL API -> OSS -> PDF/DOCX -> real Document Intelligence -> Policy/RAG/Rerank -> real Review -> Quality Gate -> persisted report; Facts reuse, new Policy/index and historical isolation: PASSED",
                flush=True,
            )


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--models-only",
        action="store_true",
        help="Only real Embedding/Rerank calls; no DB/OSS acceptance claimed",
    )
    parser.add_argument(
        "--review",
        action="store_true",
        help="Also run real HTTP upload, Document Intelligence, Review and Quality Gate with fixture-only DB factories",
    )
    args = parser.parse_args()
    if args.models_only:
        model_check()
        return
    database_name = f"contract_knowledge_test_{uuid4().hex[:12]}"
    with psycopg.connect(
        _psycopg_url(POSTGRES_DATABASE), autocommit=True, connect_timeout=5
    ) as admin:
        admin.execute(
            sql.SQL("CREATE DATABASE {}").format(sql.Identifier(database_name))
        )
    engine = create_engine(_url(database_name), connect_args={"connect_timeout": 5})
    try:
        User.__table__.create(engine)
        config = Config(str(BACKEND_ROOT / "alembic.ini"))
        config.set_main_option(
            "sqlalchemy.url", _url(database_name).render_as_string(hide_password=False)
        )
        command.upgrade(config, "head")
        command.check(config)
        asyncio.run(verify(database_name, review=args.review))
    finally:
        engine.dispose()
        with psycopg.connect(
            _psycopg_url(POSTGRES_DATABASE), autocommit=True, connect_timeout=5
        ) as admin:
            _drop_database(admin, database_name)
        print("Disposable test database cleanup complete", flush=True)


if __name__ == "__main__":
    main()
