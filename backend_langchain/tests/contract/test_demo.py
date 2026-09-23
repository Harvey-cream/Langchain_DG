import unittest
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZipFile
from unittest.mock import patch, AsyncMock
from infrastructure.document.contract_parser import validate_file, parse_contract


class _Transaction:
    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return False


class _Session:
    def begin(self):
        return _Transaction()

    def in_transaction(self):
        return False


class ParserTests(unittest.TestCase):
    def test_rejects_empty_fake_and_unsupported_uploads(self):
        for name, data in [('x.pdf', b''), ('x.pdf', b'not a pdf'), ('x.docx', b'PKfake'), ('x.exe', b'%PDF-')]:
            with self.subTest(name=name, data=data), self.assertRaises(ValueError):
                validate_file(name, data)

    def test_docx_keeps_paragraph_and_table_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph('采购合同')
        document.add_table(rows=1, cols=1).cell(0, 0).text = '金额：100 元'
        document.add_paragraph('期限：一年')
        stream = BytesIO(); document.save(stream)
        text = parse_contract('test.docx', stream.getvalue())
        self.assertLess(text.index('采购合同'), text.index('金额'))
        self.assertLess(text.index('金额'), text.index('期限'))

    def test_scan_pdf_is_not_reported_as_success(self):
        from fpdf import FPDF
        pdf = FPDF(); pdf.add_page()
        with self.assertRaisesRegex(ValueError, '扫描件'):
            parse_contract('scan.pdf', bytes(pdf.output()))

    def test_zip_without_document_is_rejected(self):
        stream = BytesIO()
        with ZipFile(stream, 'w') as archive:
            archive.writestr('random.txt', 'hello')
        with self.assertRaises(ValueError):
            validate_file('file.docx', stream.getvalue())

    def test_long_contract_is_rejected_instead_of_truncated(self):
        from docx import Document
        document = Document(); document.add_paragraph('A' * 80001)
        stream = BytesIO(); document.save(stream)
        with self.assertRaisesRegex(ValueError, '未进行截断'):
            parse_contract('long.docx', stream.getvalue())


class ReviewTests(unittest.IsolatedAsyncioTestCase):
    async def test_fabricated_evidence_is_rejected(self):
        from products.contract.application.analysis import review_document
        from products.contract.schemas.analysis import ContractAnalysis, Risk
        result = ContractAnalysis(document_type='合同', summary='摘要', risks=[Risk(
            title='风险', risk_level='high', original_text='不存在的条款', reason='原因', suggestion='建议')])
        with self.assertRaisesRegex(ValueError, '引用原文'):
            await review_document('真实条款', AsyncMock(return_value=result))

    async def test_valid_evidence_is_accepted(self):
        from products.contract.application.analysis import review_document
        from products.contract.schemas.analysis import ContractAnalysis, Risk
        result = ContractAnalysis(document_type='合同', summary='摘要', risks=[Risk(
            title='风险', risk_level='high', original_text='真实条款', reason='原因', suggestion='建议')])
        self.assertEqual(await review_document('真实 条款', AsyncMock(return_value=result)), result)

    async def test_clause_evidence_must_exist(self):
        from products.contract.application.analysis import review_document
        from products.contract.schemas.analysis import Clause, ContractAnalysis
        result = ContractAnalysis(
            document_type='合同',
            summary='摘要',
            clauses=[Clause(title='付款', original_text='编造的付款条款')],
        )
        with self.assertRaisesRegex(ValueError, '引用原文'):
            await review_document('真实合同条款', AsyncMock(return_value=result))

    async def test_risk_clause_sequence_must_reference_a_clause(self):
        from products.contract.application.analysis import review_document
        from products.contract.schemas.analysis import Clause, ContractAnalysis, Risk
        result = ContractAnalysis(
            document_type='合同',
            summary='摘要',
            clauses=[Clause(title='付款', original_text='付款条款')],
            risks=[Risk(
                title='风险', risk_level='high', original_text='付款条款',
                reason='原因', suggestion='建议', clause_sequence=2,
            )],
        )
        with self.assertRaisesRegex(ValueError, '引用原文'):
            await review_document('付款条款', AsyncMock(return_value=result))


class WorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def test_cached_content_skips_parser_and_completes(self):
        from products.contract.domain.review import ReviewContext, VersionContent
        from products.contract.schemas.analysis import ContractAnalysis
        from products.contract.workflows.review_workflow import ContractReviewWorkflow

        context = ReviewContext(uuid4(), uuid4(), 'key', 'contract.docx')
        content = VersionContent(context.version_id, '真实合同条款', 'hash', 'python-docx')
        result = ContractAnalysis(document_type='合同', summary='摘要')
        store = SimpleNamespace(
            claim_run=AsyncMock(return_value=context),
            get_content=AsyncMock(return_value=content),
            mark_parsing=AsyncMock(),
            save_content=AsyncMock(),
            mark_analyzing=AsyncMock(),
            complete_run=AsyncMock(),
            fail_run=AsyncMock(),
        )
        parser = SimpleNamespace(parse=AsyncMock())
        reviewer = SimpleNamespace(review=AsyncMock(return_value=result))

        await ContractReviewWorkflow(_Session(), store, parser, reviewer).run(context.run_id)

        parser.parse.assert_not_awaited()
        store.save_content.assert_not_awaited()
        store.mark_analyzing.assert_awaited_once_with(context)
        store.complete_run.assert_awaited_once_with(context, result)
        store.fail_run.assert_not_awaited()

    async def test_new_content_is_saved_before_review(self):
        from products.contract.domain.review import ReviewContext, VersionContent
        from products.contract.schemas.analysis import ContractAnalysis
        from products.contract.workflows.review_workflow import ContractReviewWorkflow

        context = ReviewContext(uuid4(), uuid4(), 'key', 'contract.pdf')
        content = VersionContent(context.version_id, '合同文字', 'hash', 'pdfplumber')
        store = SimpleNamespace(
            claim_run=AsyncMock(return_value=context),
            get_content=AsyncMock(return_value=None),
            mark_parsing=AsyncMock(),
            save_content=AsyncMock(),
            mark_analyzing=AsyncMock(),
            complete_run=AsyncMock(),
            fail_run=AsyncMock(),
        )
        parser = SimpleNamespace(parse=AsyncMock(return_value=content))
        reviewer = SimpleNamespace(
            review=AsyncMock(return_value=ContractAnalysis(document_type='合同', summary='摘要'))
        )

        await ContractReviewWorkflow(_Session(), store, parser, reviewer).run(context.run_id)

        parser.parse.assert_awaited_once_with(context)
        store.mark_parsing.assert_awaited_once_with(context)
        store.save_content.assert_awaited_once_with(content)

    async def test_failure_is_persisted(self):
        from products.contract.domain.review import ReviewContext, VersionContent
        from products.contract.workflows.review_workflow import ContractReviewWorkflow

        context = ReviewContext(uuid4(), uuid4(), 'key', 'contract.docx')
        content = VersionContent(context.version_id, '合同文字', 'hash', 'python-docx')
        store = SimpleNamespace(
            claim_run=AsyncMock(return_value=context),
            get_content=AsyncMock(return_value=content),
            mark_parsing=AsyncMock(),
            save_content=AsyncMock(),
            mark_analyzing=AsyncMock(),
            complete_run=AsyncMock(),
            fail_run=AsyncMock(),
        )
        parser = SimpleNamespace(parse=AsyncMock())
        reviewer = SimpleNamespace(review=AsyncMock(side_effect=RuntimeError('provider down')))

        await ContractReviewWorkflow(_Session(), store, parser, reviewer).run(context.run_id)

        store.complete_run.assert_not_awaited()
        store.fail_run.assert_awaited_once_with(
            context.run_id,
            '分析服务暂不可用或超时，请稍后重新分析',
        )


if __name__ == '__main__':
    unittest.main()
