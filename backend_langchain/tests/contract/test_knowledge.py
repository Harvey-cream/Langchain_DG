"""Real SQLite/Parser/API tests, explicitly fake external Embedding/Rerank/OSS."""

import asyncio
import math
from io import BytesIO
from types import SimpleNamespace
from uuid import UUID, uuid4
from unittest.mock import AsyncMock

import httpx
import pytest
from docx import Document as Docx
from fastapi import FastAPI
from sqlalchemy import func, select

from common.account.models import User
from common.auth.dependencies import require_user
from common.database import get_db
from products.contract import api, db, retrieval
from products.contract.document.contract_parser import parse_contract
from products.contract.knowledge import (
    chunk_document,
    ingest_document,
    finish_document_delete,
)
from products.contract.models import (
    ContractKnowledgeDocumentModel as Document,
    ContractKnowledgeChunkModel as Chunk,
)


class MemoryStorage:
    def __init__(self):
        self.objects = {}
        self.fail_put = self.fail_delete = False

    def put(self, key, data, **kwargs):
        if self.fail_put:
            raise RuntimeError("fake upload failure")
        self.objects[key] = data

    def download(self, key, path):
        path.write_bytes(self.objects[key])

    def delete(self, key):
        if self.fail_delete:
            raise RuntimeError("fake delete failure")
        self.objects.pop(key, None)


class FakeEmbeddings:
    def __init__(self):
        self.calls = []

    def embed_documents(self, texts):
        self.calls.append(list(texts))
        return [[1.0, 1.0] + [0.0] * 1022 for _ in texts]

    def embed_query(self, query):
        self.calls.append(query)
        return [1.0, 1.0] + [0.0] * 1022


def fake_rerank(query, texts, *, top_n, timeout):
    return list(reversed(range(len(texts))))[:top_n]


def docx_bytes(text="企业采购规范：付款期限不超过45天。"):
    document, stream = Docx(), BytesIO()
    document.add_paragraph(text)
    document.save(stream)
    return stream.getvalue()


async def seed_document(session, storage, data, *, owner_id=1, filename="参考.docx"):
    key = f"knowledge-test/{uuid4()}"
    document = Document(
        user_id=owner_id,
        filename=filename,
        source_key=key,
        content_hash=uuid4().hex * 2,
        byte_count=len(data),
        status="uploaded",
    )
    session.add(document)
    await session.flush()
    storage.put(key, data)
    return document


def test_actual_bm25_and_rrf_math():
    scores = retrieval.bm25_scores("payment", [["payment"], ["delivery"]])
    assert math.isclose(scores[0], math.log(2)) and scores[1] == 0
    assert retrieval.rrf(["a", "b"], ["b", "c"]) == ["b", "a", "c"]
    assert retrieval.rrf(["a", "a"], ["b"]) == ["a", "b"]
    assert "付款" in retrieval.tokenize("付款期限 payment PAYMENT")


@pytest.mark.parametrize("kind", ["docx", "pdf"])
def test_pdf_docx_retrieval_chunking_preserves_source_and_real_locator(kind):
    if kind == "docx":
        data = docx_bytes("付款期限" * 4000)
    else:
        from fpdf import FPDF

        pdf = FPDF()
        for text in ("Payment standard", "Delivery example"):
            pdf.add_page()
            pdf.set_font("Helvetica", size=12)
            pdf.cell(text=text)
        data = bytes(pdf.output())
    parsed = parse_contract(f"reference.{kind}", data)
    chunks = chunk_document(parsed)
    from products.contract.domain.document import join_block_text

    text = join_block_text(parsed.blocks)
    for chunk in chunks:
        locator = chunk["locator"]
        assert chunk["chunk_text"] == text[locator["start_char"] : locator["end_char"]]
        assert len(chunk["chunk_text"]) <= 1200
        assert locator["block_sequences"]
    reconstructed = chunks[0]["chunk_text"] + "".join(
        c["chunk_text"][150:] for c in chunks[1:]
    )
    assert reconstructed == text
    assert (
        (chunks[0]["locator"]["pages"] == [1, 2])
        if kind == "pdf"
        else not chunks[0]["locator"]["pages"]
    )


def test_ingestion_hybrid_isolation_idempotence_failure_and_reindex(document_db):
    factory, _ = document_db
    storage, embedding = MemoryStorage(), FakeEmbeddings()

    async def check():
        async with factory() as session, session.begin():
            session.add(
                User(
                    user_id=2,
                    email="private@test.invalid",
                    username="private",
                    password="fake",
                )
            )
            await session.flush()
            first = await seed_document(session, storage, docx_bytes())
            private = await seed_document(
                session, storage, docx_bytes("用户B绝密付款规则。"), owner_id=2
            )
            first_id, private_id = first.id, private.id
        await asyncio.gather(
            ingest_document(
                first_id, 1, factory=factory, storage=storage, embeddings=embedding
            ),
            ingest_document(
                first_id, 1, factory=factory, storage=storage, embeddings=embedding
            ),
        )
        await ingest_document(
            private_id, 1, factory=factory, storage=storage, embeddings=embedding
        )
        assert len(embedding.calls) == 2
        rankings = []

        def rerank(query, texts, **kwargs):
            rankings.append(texts)
            return fake_rerank(query, texts, **kwargs)

        references = await retrieval.retrieve_references(
            factory, 1, "付款期限规范", embeddings=embedding, reranker=rerank
        )
        assert references and all(
            r["owner_id"] == 1 and r["document_id"] == str(first_id) for r in references
        )
        assert not any("绝密" in text for text in rankings[-1])
        with pytest.raises(ValueError):
            await retrieval.retrieve_references(
                factory, 0, "付款", embeddings=embedding, reranker=rerank
            )
        empty_model = SimpleNamespace(
            embed_query=lambda q: (_ for _ in ()).throw(
                AssertionError("No model needed")
            )
        )
        assert (
            await retrieval.retrieve_references(
                factory, 999, "付款", embeddings=empty_model, reranker=rerank
            )
            == []
        )
        async with factory() as session, session.begin():
            stored = await session.get(Document, first_id)
            old_ids = list(
                (
                    await session.scalars(
                        select(Chunk.id).where(Chunk.document_id == first_id)
                    )
                ).all()
            )
            stored.generation, stored.status = 2, "uploaded"
        bad = SimpleNamespace(embed_documents=lambda texts: [[1, 2] for _ in texts])
        await ingest_document(
            first_id, 2, factory=factory, storage=storage, embeddings=bad
        )
        async with factory() as session, session.begin():
            stored = await session.get(Document, first_id)
            assert stored.status == "failed" and "维度" in stored.error
            assert (
                list(
                    (
                        await session.scalars(
                            select(Chunk.id).where(Chunk.document_id == first_id)
                        )
                    ).all()
                )
                == old_ids
            )
            stored.generation, stored.status = 3, "uploaded"
        assert (
            await retrieval.retrieve_references(
                factory, 1, "付款", embeddings=embedding, reranker=rerank
            )
            == []
        )
        await ingest_document(
            first_id, 3, factory=factory, storage=storage, embeddings=embedding
        )
        await ingest_document(
            first_id, 3, factory=factory, storage=storage, embeddings=embedding
        )
        async with factory() as session, session.begin():
            rows = (
                await session.scalars(
                    select(Chunk).where(Chunk.document_id == first_id)
                )
            ).all()
            assert len(rows) == len(old_ids) and all(
                r.generation == 3 and r.id not in old_ids for r in rows
            )
        with pytest.raises(ValueError, match="不可用"):
            await retrieval.retrieve_references(
                factory,
                1,
                "付款",
                embeddings=embedding,
                reranker=lambda *a, **k: (_ for _ in ()).throw(
                    RuntimeError("Rerank down")
                ),
            )
        slow = SimpleNamespace(
            embed_query=lambda q: (
                __import__("time").sleep(0.1),
                [1.0, 1.0] + [0.0] * 1022,
            )[1]
        )
        config = retrieval.settings() | {"timeout": 0.01}
        with pytest.raises(ValueError, match="超时"):
            await retrieval.retrieve_references(
                factory, 1, "付款", embeddings=slow, reranker=rerank, config=config
            )

    asyncio.run(check())


def test_restart_exposes_interrupted_upload_index_and_delete_as_recoverable(
    document_db, monkeypatch
):
    from products.contract import startup

    factory, _ = document_db
    monkeypatch.setattr(startup, "engine", factory.kw["bind"])

    async def check():
        ids = {}
        async with factory() as session, session.begin():
            for status in ("uploading", "uploaded", "processing", "ready", "deleting"):
                document = await seed_document(session, MemoryStorage(), docx_bytes())
                document.status = status
                ids[status] = document.id
        await startup.recover_interrupted_contract_reviews()
        async with factory() as session:
            for status, document_id in ids.items():
                document = await session.get(Document, document_id)
                if status in ("uploading", "uploaded", "processing"):
                    assert (
                        document.status == "failed"
                        and document.generation == 2
                        and document.error
                    )
                elif status == "deleting":
                    assert (
                        document.status == "deleting" and "重试删除" in document.error
                    )
                else:
                    assert document.status == "ready" and document.generation == 1

    asyncio.run(check())


def test_knowledge_upload_crud_download_dedup_and_cleanup_recovery(
    document_db, monkeypatch
):
    factory, _ = document_db
    storage, jobs, caller = MemoryStorage(), [], [1]
    monkeypatch.setattr(api, "OssObjectStorage", lambda: storage)

    async def enqueue(*args):
        jobs.append(args)

    async def cleanup(document_id, generation):
        await finish_document_delete(
            document_id, generation, factory=factory, storage=storage
        )

    monkeypatch.setattr(api, "ingest_document", enqueue)
    monkeypatch.setattr(api, "finish_document_delete", cleanup)
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[require_user] = lambda: SimpleNamespace(user_id=caller[0])

    async def sessions():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_db] = sessions

    async def check():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app), base_url="http://test"
        ) as client:
            data = docx_bytes()
            upload = await client.post(
                "/api/contract-knowledge/documents/upload",
                files={"file": ("参考.docx", data)},
                data={"category": "legal_policy"},
            )
            assert upload.status_code == 200
            document = upload.json()["data"]["document"]
            did = document["id"]
            duplicate = await client.post(
                "/api/contract-knowledge/documents/upload",
                files={"file": ("copy.docx", data)},
            )
            assert (
                duplicate.json()["data"]["duplicate"]
                and duplicate.json()["data"]["document"]["id"] == did
            )
            assert len(jobs) == 1 and len(storage.objects) == 1
            await ingest_document(
                UUID(did),
                1,
                factory=factory,
                storage=storage,
                embeddings=FakeEmbeddings(),
            )
            async with factory() as session:
                chunk_id = await session.scalar(
                    select(Chunk.id).where(Chunk.document_id == UUID(did))
                )
            caller[0] = 999
            assert (await client.get("/api/contract-knowledge/documents")).json()[
                "data"
            ]["documents"] == []
            for method, path, kwargs in [
                ("GET", f"/documents/{did}", {}),
                ("GET", f"/documents/{did}/download", {}),
                ("GET", f"/chunks/{chunk_id}", {}),
                ("DELETE", f"/documents/{did}", {}),
                ("POST", f"/documents/{did}/reindex", {}),
                (
                    "PATCH",
                    f"/documents/{did}",
                    {"json": {"category": "other", "expected_generation": 1}},
                ),
            ]:
                assert (
                    await client.request(
                        method, "/api/contract-knowledge" + path, **kwargs
                    )
                ).status_code == 404
            caller[0] = 1
            assert (
                await client.get(f"/api/contract-knowledge/documents/{did}/download")
            ).content == data
            assert (
                await client.get(f"/api/contract-knowledge/chunks/{chunk_id}")
            ).json()["data"]["document_id"] == did
            changed = await client.patch(
                f"/api/contract-knowledge/documents/{did}",
                json={"category": "custom_reference", "expected_generation": 1},
            )
            assert changed.json()["data"]["category"] == "custom_reference"
            queued = await client.post(
                f"/api/contract-knowledge/documents/{did}/reindex"
            )
            assert queued.json()["data"]["generation"] == 2
            await client.post(f"/api/contract-knowledge/documents/{did}/reindex")
            assert len(jobs) == 2
            storage.fail_delete = True
            assert (
                await client.delete(f"/api/contract-knowledge/documents/{did}")
            ).status_code == 200
            detail = (
                await client.get(f"/api/contract-knowledge/documents/{did}")
            ).json()["data"]
            assert detail["status"] == "deleting" and detail["error"]
            assert (
                await client.get(f"/api/contract-knowledge/chunks/{chunk_id}")
            ).status_code == 404
            storage.fail_delete = False
            await client.delete(f"/api/contract-knowledge/documents/{did}")
            assert (await client.get("/api/contract-knowledge/documents")).json()[
                "data"
            ]["documents"] == []
            restored = await client.post(
                "/api/contract-knowledge/documents/upload",
                files={"file": ("new.docx", data)},
            )
            assert restored.json()["data"]["document"]["id"] != did
            storage.fail_put = True
            failed_data = docx_bytes("另一个规范。")
            assert (
                await client.post(
                    "/api/contract-knowledge/documents/upload",
                    files={"file": ("failed.docx", failed_data)},
                )
            ).status_code == 503
            failed = next(
                d
                for d in (await client.get("/api/contract-knowledge/documents")).json()[
                    "data"
                ]["documents"]
                if d["status"] == "failed"
            )
            storage.fail_put = False
            recovered = await client.post(
                "/api/contract-knowledge/documents/upload",
                files={"file": ("failed.docx", failed_data)},
            )
            assert recovered.json()["data"]["document"]["id"] == failed["id"]
            assert (
                recovered.json()["data"]["document"]["generation"]
                == failed["generation"] + 1
            )

    asyncio.run(check())


def test_dense_sql_scopes_both_owner_tables_before_execution():
    from sqlalchemy.dialects import postgresql

    sql = str(
        retrieval.scoped_chunks(7)
        .with_only_columns(Chunk.id)
        .order_by(Chunk.embedding.cosine_distance([1.0] * 1024))
        .compile(dialect=postgresql.dialect())
    )
    assert "<=>" in sql and "contract_knowledge_chunks.user_id =" in sql
    assert (
        "contract_knowledge_documents.user_id =" in sql
        and "contract_knowledge_documents.status =" in sql
    )


def test_reference_fabrication_and_wrong_clause_are_rejected():
    from products.contract.domain.document import CanonicalContractClause
    from products.contract.schemas.analysis import ReviewResult, validate_findings

    vid, cid, chunk_id = uuid4(), uuid4(), uuid4()
    clause = CanonicalContractClause(
        cid, vid, 1, "付款", "payment", "付款90天。", (uuid4(),)
    )
    reference = dict(
        chunk_id=str(chunk_id), chunk_text="企业付款不超过45天。", clause_ids=[str(cid)]
    )
    risk = dict(
        clause_id=str(cid),
        risk_type="payment",
        risk_level="medium",
        evidence="付款90天。",
        reason="超过参考期限",
        recommendation="复核付款期限",
        reference_citations=[dict(chunk_id=str(chunk_id), source_excerpt="不超过45天")],
    )
    assert validate_findings(
        ReviewResult.model_validate({"risks": [risk]}),
        (clause,),
        vid,
        references=[reference],
    )
    for citation in [
        dict(chunk_id=str(uuid4()), source_excerpt="不超过45天"),
        dict(chunk_id=str(chunk_id), source_excerpt="编造法规"),
    ]:
        with pytest.raises(ValueError):
            validate_findings(
                ReviewResult.model_validate(
                    {"risks": [risk | {"reference_citations": [citation]}]}
                ),
                (clause,),
                vid,
                references=[reference],
            )
    with pytest.raises(ValueError):
        validate_findings(
            ReviewResult.model_validate({"risks": [risk]}),
            (clause,),
            vid,
            references=[reference | {"clause_ids": []}],
        )


def test_review_dynamic_retrieval_and_history_survive_reindex_and_delete(
    document_db, monkeypatch
):
    import json
    from products.contract.agents.review import ReviewAgent, ReviewLimits
    from products.contract.domain.document import (
        DocumentBlock,
        ParsedDocument,
        DocumentFacts,
        CanonicalContractClause,
    )
    from products.contract.document_intelligence.validation import validate_clauses
    from products.contract.document_intelligence.pipeline import (
        DocumentIntelligencePipeline,
    )
    from products.contract.models import (
        AnalysisRunModel,
        ContractPolicyModel,
        ContractVersionModel,
        ContractRiskModel,
    )
    from products.contract.workflows.review_workflow import run_contract_review

    factory, vid = document_db
    storage, embedding, model_calls = MemoryStorage(), FakeEmbeddings(), []
    monkeypatch.setattr(
        retrieval, "make_dashscope_embeddings", lambda **kwargs: embedding
    )
    monkeypatch.setattr(retrieval, "rerank_indices", fake_rerank)

    async def invoke(messages):
        payload = json.loads(messages[1].content)
        model_calls.append(payload)
        target, policy, reference = (
            payload["targets"][0],
            payload["policies"][0],
            payload["references"][0],
        )
        return {
            "risks": [
                dict(
                    clause_id=target["clause_id"],
                    risk_type="payment",
                    risk_level="high",
                    evidence=target["original_text"],
                    reason="不符合企业规则，参考资料提供协商案例",
                    recommendation="业务确认后协商",
                    policy_ids=[policy["policy_id"]],
                    reference_citations=[
                        dict(
                            chunk_id=reference["chunk_id"],
                            source_excerpt=reference["chunk_text"][:25],
                        )
                    ],
                )
            ]
        }

    model = SimpleNamespace(
        with_structured_output=lambda *args, **kwargs: SimpleNamespace(ainvoke=invoke)
    )

    async def check():
        async with factory() as session, session.begin():
            document = await seed_document(session, storage, docx_bytes())
            did = document.id
            block = DocumentBlock(uuid4(), vid, 1, "甲方付款期限90天。")
            clause = CanonicalContractClause(
                uuid4(), vid, 1, "付款", "payment", block.text, (block.block_id,)
            )
            parsed = ParsedDocument(vid, block.text, (block,), None, "python-docx")
            facts = DocumentFacts(
                parsed,
                (clause,),
                {
                    "summary": "固定的版本事实",
                    "parties": [],
                    "amount": "",
                    "duration": "",
                    "payment_terms": "",
                },
                validate_clauses(vid, (block,), (clause,)),
            )
            await db.save_document_facts(session, facts)
            policy = ContractPolicyModel(
                user_id=1, name="期限", requirement="付款不超过45天", severity="high"
            )
            session.add(policy)
            await session.flush()
            pid = policy.id
            contract_id = (await session.get(ContractVersionModel, vid)).contract_id
        await ingest_document(
            did, 1, factory=factory, storage=storage, embeddings=embedding
        )
        run_ids, snapshots = [], []
        for attempt in (1, 2):
            async with factory() as session:
                async with session.begin():
                    run = AnalysisRunModel(
                        version_id=vid,
                        attempt=attempt,
                        status="pending",
                        our_party="甲方",
                        party_role="buyer",
                        contract_type="procurement",
                    )
                    session.add(run)
                    await session.flush()
                    run_ids.append(run.id)
                parser = SimpleNamespace(
                    parse=AsyncMock(
                        side_effect=AssertionError("Review retry must not reparse")
                    )
                )
                pipeline = DocumentIntelligencePipeline(
                    session, parser, SimpleNamespace()
                )
                await run_contract_review(
                    session,
                    pipeline,
                    ReviewAgent(
                        model,
                        limits=ReviewLimits(react_enabled=False),
                        retrieval_factory=factory,
                    ),
                    run.id,
                )
                async with session.begin():
                    await session.refresh(run)
                    assert run.status == "completed", run.error
                    snapshots.append(run.context_snapshot)
                    assert await db.get_document_facts(session, vid) == facts
                    risk = await session.scalar(
                        select(ContractRiskModel).where(
                            ContractRiskModel.run_id == run.id
                        )
                    )
                    assert risk.policy_ids == [str(pid)] and risk.reference_citations
                    assert parser.parse.await_count == 0
            if attempt == 1:
                async with factory() as session, session.begin():
                    rule = await session.get(ContractPolicyModel, pid)
                    rule.version, rule.requirement = 2, "付款不超过30天"
                    source = await session.get(Document, did)
                    source.status, source.generation = "uploaded", 2
                await ingest_document(
                    did, 2, factory=factory, storage=storage, embeddings=embedding
                )
        assert len(model_calls) == 2
        assert (
            snapshots[0]["policies"][0]["version"] == 1
            and snapshots[1]["policies"][0]["version"] == 2
        )
        assert (
            snapshots[0]["references"][0]["index_generation"] == 1
            and snapshots[1]["references"][0]["index_generation"] == 2
        )
        assert (
            snapshots[0]["references"][0]["chunk_id"]
            != snapshots[1]["references"][0]["chunk_id"]
        )
        async with factory() as session, session.begin():
            from sqlalchemy import delete

            source = await session.get(Document, did)
            source.status, source.generation = "deleting", 3
            await session.execute(delete(Chunk).where(Chunk.document_id == did))
        await finish_document_delete(did, 3, factory=factory, storage=storage)
        async with factory() as session, session.begin():
            assert (
                await session.get(AnalysisRunModel, run_ids[0])
            ).context_snapshot == snapshots[0]
            assert await db.select_run(session, 1, contract_id, vid, run_ids[0])
            selected = await db.load_snapshot(session, 1, contract_id, vid)
            assert (
                selected.selected_run.context_snapshot == snapshots[0]
                and selected.facts == facts
            )
        app = FastAPI()
        app.include_router(api.router)
        caller = [1]
        app.dependency_overrides[require_user] = lambda: SimpleNamespace(
            user_id=caller[0]
        )

        async def sessions():
            async with factory() as session:
                yield session

        app.dependency_overrides[get_db] = sessions
        path = f'/api/contracts/{contract_id}/versions/{vid}/analysis-runs/{run_ids[0]}/references/{snapshots[0]["references"][0]["chunk_id"]}'
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app), base_url="http://test"
        ) as client:
            assert (await client.get(path)).json()["data"]["chunk_text"] == snapshots[
                0
            ]["references"][0]["chunk_text"]
            caller[0] = 999
            assert (await client.get(path)).status_code == 404

    asyncio.run(check())


def test_deleted_document_cannot_be_resurrected_by_late_ingestion(document_db):
    import threading
    from sqlalchemy import delete

    factory, _ = document_db
    storage, entered, released = MemoryStorage(), threading.Event(), threading.Event()

    class SlowEmbedding(FakeEmbeddings):
        def embed_documents(self, texts):
            entered.set()
            if not released.wait(timeout=5):
                raise RuntimeError("test barrier timeout")
            return super().embed_documents(texts)

    async def check():
        async with factory() as session, session.begin():
            document = await seed_document(session, storage, docx_bytes())
            did = document.id
        task = asyncio.create_task(
            ingest_document(
                did, 1, factory=factory, storage=storage, embeddings=SlowEmbedding()
            )
        )
        try:
            assert await asyncio.to_thread(entered.wait, 5)
            async with factory() as session, session.begin():
                document = await session.get(Document, did)
                document.status, document.generation = "deleting", 2
                await session.execute(delete(Chunk).where(Chunk.document_id == did))
            await finish_document_delete(did, 2, factory=factory, storage=storage)
        finally:
            released.set()
        await task
        async with factory() as session:
            assert (await session.get(Document, did)).status == "deleted"
            assert await session.scalar(select(func.count()).select_from(Chunk)) == 0
            assert not storage.objects

    asyncio.run(check())


def test_model_rerank_transport_uses_provider_scores_and_rejects_invalid_output(
    monkeypatch,
):
    from common.integrations import retrieval_models

    requests = []

    def reply(request):
        requests.append(__import__("json").loads(request.content))
        return httpx.Response(
            200,
            json={
                "output": {
                    "results": [
                        {"index": 1, "relevance_score": 0.9},
                        {"index": 0, "relevance_score": 0.1},
                    ]
                }
            },
        )

    original_client = httpx.Client
    monkeypatch.setattr(retrieval_models, "dashscope_api_key", lambda: "fake-key")
    monkeypatch.setattr(
        retrieval_models.httpx,
        "Client",
        lambda **kwargs: original_client(transport=httpx.MockTransport(reply)),
    )
    assert retrieval_models.rerank_indices(
        "付款", ["交付约定", "付款案例"], top_n=2
    ) == [1, 0]
    assert requests[0]["input"]["documents"] == ["交付约定", "付款案例"]
    monkeypatch.setattr(retrieval_models, "dashscope_api_key", lambda: "")
    with pytest.raises(RuntimeError):
        retrieval_models.rerank_indices("付款", ["付款"], top_n=1)


def test_index_publication_rollback_keeps_old_rows_unsearchable(document_db):
    from sqlalchemy import event

    factory, _ = document_db
    storage, embedding = MemoryStorage(), FakeEmbeddings()

    async def check():
        async with factory() as session, session.begin():
            document = await seed_document(session, storage, docx_bytes())
            did = document.id
        await ingest_document(
            did, 1, factory=factory, storage=storage, embeddings=embedding
        )
        async with factory() as session, session.begin():
            ids = list((await session.scalars(select(Chunk.id))).all())
            document = await session.get(Document, did)
            document.status, document.generation = "uploaded", 2
        engine = factory.kw["bind"].sync_engine

        def fail_after_delete(connection, cursor, statement, parameters, context, many):
            if (
                statement.lstrip()
                .upper()
                .startswith("INSERT INTO CONTRACT_KNOWLEDGE_CHUNKS")
            ):
                raise RuntimeError("fake vector index publication failure")

        event.listen(engine, "before_cursor_execute", fail_after_delete)
        try:
            await ingest_document(
                did, 2, factory=factory, storage=storage, embeddings=embedding
            )
        finally:
            event.remove(engine, "before_cursor_execute", fail_after_delete)
        async with factory() as session:
            document = await session.get(Document, did)
            assert document.status == "failed" and document.error
            assert list((await session.scalars(select(Chunk.id))).all()) == ids
        assert (
            await retrieval.retrieve_references(
                factory, 1, "付款", embeddings=embedding, reranker=fake_rerank
            )
            == []
        )

    asyncio.run(check())


@pytest.mark.parametrize(
    "fail_put", [False, True], ids=["late-success", "late-failure"]
)
def test_delete_during_upload_cleans_late_oss_object(
    document_db, monkeypatch, fail_put
):
    import threading

    factory, _ = document_db
    entered, release = threading.Event(), threading.Event()

    class SlowStorage(MemoryStorage):
        def put(self, key, data, **kwargs):
            entered.set()
            if not release.wait(timeout=5):
                raise RuntimeError("test barrier timeout")
            super().put(key, data, **kwargs)

    storage = SlowStorage()
    storage.fail_put = fail_put
    monkeypatch.setattr(api, "OssObjectStorage", lambda: storage)

    async def cleanup(did, generation):
        await finish_document_delete(did, generation, factory=factory, storage=storage)

    monkeypatch.setattr(api, "finish_document_delete", cleanup)
    app = FastAPI()
    app.include_router(api.router)
    app.dependency_overrides[require_user] = lambda: SimpleNamespace(user_id=1)

    async def sessions():
        async with factory() as session:
            yield session

    app.dependency_overrides[get_db] = sessions

    async def check():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app), base_url="http://test"
        ) as client:
            upload = asyncio.create_task(
                client.post(
                    "/api/contract-knowledge/documents/upload",
                    files={"file": ("race.docx", docx_bytes())},
                )
            )
            try:
                assert await asyncio.to_thread(entered.wait, 5)
                document = (
                    await client.get("/api/contract-knowledge/documents")
                ).json()["data"]["documents"][0]
                assert document["status"] == "uploading"
                assert (
                    await client.delete(
                        f'/api/contract-knowledge/documents/{document["id"]}'
                    )
                ).status_code == 200
            finally:
                release.set()
            assert (await upload).status_code == (503 if fail_put else 409)
            assert not storage.objects
            async with factory() as session:
                assert (
                    await session.get(Document, UUID(document["id"]))
                ).status == "deleted"

    asyncio.run(check())
