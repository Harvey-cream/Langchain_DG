import asyncio
from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from docx import Document
from fpdf import FPDF
from sqlalchemy import func, select

from products.contract.document.contract_parser import parse_contract
from products.contract.models import (
    AnalysisRunModel,
    CanonicalContractClauseModel,
    ContractClauseModel,
    ContractDocumentBlockModel,
    ContractDocumentFactsModel,
    ContractRiskModel,
    ContractVersionModel,
)
from products.contract import db as contract_db
from products.contract.document_intelligence.pipeline import DocumentIntelligencePipeline
from products.contract.domain.review import ReviewContext, VersionContent
from products.contract.schemas.analysis import Clause, ContractAnalysis, Risk
from products.contract.schemas.document import ClauseExtraction, DocumentMetadata, SectionExtraction
from products.contract.workflows.review_workflow import run_contract_review


def fixture_file(kind):
    stream = BytesIO()
    if kind == 'docx':
        doc = Document()
        doc.add_paragraph('第一条 付款')
        doc.add_paragraph('甲方支付100元。')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '期限：一年'
        doc.save(stream)
        return stream.getvalue()
    pdf = FPDF()
    for text in ['1. Payment', 'Pay 100 USD.']:
        pdf.add_page()
        pdf.set_font('Helvetica', size=12)
        pdf.cell(text=text)
    return bytes(pdf.output())


def capabilities(session):
    async def extract(blocks):
        assert not session.in_transaction(), 'LLM must run outside DB transaction'
        return ClauseExtraction.model_validate({'clauses': [{
            'title': '条款', 'clause_type': 'payment', 'source_block_ids': [str(b.block_id) for b in blocks],
        }]})

    async def sections(blocks):
        assert not session.in_transaction()
        return SectionExtraction.model_validate({'sections': [{
            'title': '正文', 'source_block_ids': [str(b.block_id) for b in blocks],
        }]})

    async def metadata(blocks):
        assert not session.in_transaction()
        return DocumentMetadata(summary='合同事实', parties=['甲方'], amount='100', duration='一年')

    return SimpleNamespace(clauses=AsyncMock(side_effect=extract), sections=AsyncMock(side_effect=sections),
                           metadata=AsyncMock(side_effect=metadata))


@pytest.mark.parametrize('kind', ['pdf', 'docx'])
def test_full_pipeline_review_retry_history_and_facts_reused(document_db, kind):
    factory, version_id = document_db

    async def run():
        async with factory() as session:
            async def parse(context):
                assert not session.in_transaction()
                return parse_contract(f'test.{kind}', fixture_file(kind), version_id=context.version_id)
            parser = SimpleNamespace(parse=AsyncMock(side_effect=parse))
            fake = capabilities(session)
            pipeline = DocumentIntelligencePipeline(session, parser, fake)

            async def review(text):
                assert not session.in_transaction()
                return ContractAnalysis(document_type='合同', summary='Review V1 summary',
                    clauses=[Clause(title='旧版条款', original_text=text)],
                    risks=[Risk(title='Review V1 risk', risk_level='low', original_text=text,
                                reason='review reasoning', suggestion='review suggestion', clause_sequence=1)])
            reviewer = SimpleNamespace(review=AsyncMock(side_effect=review))
            run_ids = []
            first = None
            for attempt in (1, 2):
                async with session.begin():
                    row = AnalysisRunModel(
                        id=uuid4(),
                        version_id=version_id,
                        status="pending",
                        current_step="queued",
                        attempt=attempt,
                        prompt_version="v2",
                    )
                    session.add(row)
                    await session.flush()
                    run_ids.append(row.id)
                await run_contract_review(session, pipeline, reviewer, row.id)
                async with session.begin():
                    assert row.status == 'completed', row.error
                    facts = await contract_db.get_document_facts(session, version_id)
                    if first is None:
                        first = facts
                    assert facts == first
                    assert facts.coverage.ratio == 1
                    assert facts.metadata['amount'] == '100'
                    assert await session.scalar(select(func.count()).select_from(ContractDocumentBlockModel)) == len(facts.document.blocks)
                    assert await session.scalar(select(func.count()).select_from(CanonicalContractClauseModel)) == len(facts.clauses)
            parser.parse.assert_awaited_once()
            assert reviewer.review.await_count == 2
            assert fake.metadata.await_count == 1
            async with session.begin():
                version = await session.get(ContractVersionModel, version_id)
                history = await contract_db.list_history(
                    session, 1, version.contract_id, version_id
                )
                assert len(history) == 2
                assert await contract_db.select_run(
                    session, 1, version.contract_id, version_id, run_ids[0]
                )
                snapshot = await contract_db.load_snapshot(
                    session, 1, version.contract_id, version_id
                )
                assert snapshot.selected_run.id == run_ids[0]
                assert snapshot.selected_run.result['summary'] == 'Review V1 summary'
                assert snapshot.risks[0].clause_id == snapshot.clauses[0].id
                assert snapshot.document_text == facts.document.document_text
                assert await session.scalar(select(func.count()).select_from(ContractClauseModel)) == 2
    asyncio.run(run())


def test_invalid_extraction_leaves_no_partial_facts_and_marks_run_failed(document_db):
    factory, version_id = document_db

    async def run():
        async with factory() as session:
            fake = capabilities(session)
            fake.clauses = AsyncMock(return_value={'clauses': [{'title': 'fabricated', 'source_block_ids': [str(uuid4())]}]})
            parser = SimpleNamespace(parse=AsyncMock(return_value=parse_contract('test.docx', fixture_file('docx'), version_id=version_id)))
            pipeline = DocumentIntelligencePipeline(session, parser, fake)
            reviewer = SimpleNamespace(review=AsyncMock())
            async with session.begin():
                row = AnalysisRunModel(
                    id=uuid4(),
                    version_id=version_id,
                    status="pending",
                    current_step="queued",
                    attempt=1,
                    prompt_version="v2",
                )
                session.add(row)
                await session.flush()
            await run_contract_review(session, pipeline, reviewer, row.id)
            async with session.begin():
                assert row.status == 'failed'
                assert 'Document Intelligence' in row.error
                assert await contract_db.get_content(session, version_id) is None
                assert await contract_db.get_document_facts(session, version_id) is None
                assert await session.scalar(select(func.count()).select_from(ContractDocumentBlockModel)) == 0
            assert fake.clauses.await_count == 2
            reviewer.review.assert_not_awaited()
    asyncio.run(run())


def test_lazy_reparse_preserves_historical_review_text(document_db):
    factory, version_id = document_db

    async def run():
        async with factory() as session:
            async with session.begin():
                await contract_db.save_content(
                    session,
                    VersionContent(version_id, '旧版原文', 'legacy-hash', 'old', 'v1'),
                )
                old = AnalysisRunModel(
                    id=uuid4(),
                    version_id=version_id,
                    status="pending",
                    current_step="queued",
                    attempt=1,
                    prompt_version="v2",
                )
                session.add(old)
                await session.flush()
                old.status = 'completed'
                old.result = ContractAnalysis(document_type='合同', summary='old summary').model_dump()
            parser = SimpleNamespace(parse=AsyncMock(return_value=parse_contract('test.docx', fixture_file('docx'), version_id=version_id)))
            pipeline = DocumentIntelligencePipeline(session, parser, capabilities(session))
            await pipeline.ensure(ReviewContext(uuid4(), version_id, 'key', 'test.docx'))
            async with session.begin():
                version = await session.get(ContractVersionModel, version_id)
                snapshot = await contract_db.load_snapshot(
                    session, 1, version.contract_id, version_id
                )
                assert snapshot.document_text == '旧版原文'
                assert snapshot.selected_run.result['summary'] == 'old summary'
                facts = await contract_db.get_document_facts(session, version_id)
                assert facts.document.parser_version == 'v2'
    asyncio.run(run())


def test_metadata_forbids_risk_fields():
    with pytest.raises(ValueError):
        DocumentMetadata.model_validate({'risk': 'not document facts'})


def test_publish_failure_rolls_back_content_blocks_clauses_and_marker(document_db):
    factory, version_id = document_db

    async def run():
        async with factory() as session:
            save_clauses = contract_db.save_canonical_clauses
            async def fail_after_clauses(*args):
                await save_clauses(*args)
                raise RuntimeError('simulated final publication failure')
            parser = SimpleNamespace(parse=AsyncMock(return_value=parse_contract('test.docx', fixture_file('docx'), version_id=version_id)))
            pipeline = DocumentIntelligencePipeline(session, parser, capabilities(session))
            from unittest.mock import patch
            with patch.object(contract_db, 'save_canonical_clauses', fail_after_clauses):
                with pytest.raises(RuntimeError, match='publication'):
                    await pipeline.ensure(ReviewContext(uuid4(), version_id, 'key', 'test.docx'))
            async with session.begin():
                assert await contract_db.get_content(session, version_id) is None
                for model in (ContractDocumentBlockModel, CanonicalContractClauseModel, ContractDocumentFactsModel):
                    assert await session.scalar(select(func.count()).select_from(model)) == 0
    asyncio.run(run())


def test_low_coverage_is_persisted_as_diagnostic(document_db):
    factory, version_id = document_db

    async def run():
        async with factory() as session:
            fake = capabilities(session)
            async def partial(blocks):
                return {'clauses': [{'title': '部分条款', 'source_block_ids': [str(blocks[0].block_id)]}]}
            fake.clauses = AsyncMock(side_effect=partial)
            parser = SimpleNamespace(parse=AsyncMock(return_value=parse_contract('test.docx', fixture_file('docx'), version_id=version_id)))
            facts = await DocumentIntelligencePipeline(session, parser, fake).ensure(
                ReviewContext(uuid4(), version_id, 'key', 'test.docx'))
            assert facts.coverage.ratio == 1 / 3
            assert facts.coverage.warnings
            async with session.begin():
                stored = await contract_db.get_document_facts(session, version_id)
                model = await session.get(ContractDocumentFactsModel, version_id)
                assert stored.coverage == facts.coverage
                assert model.warnings == list(facts.coverage.warnings)
    asyncio.run(run())
