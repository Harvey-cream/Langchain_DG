from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from fpdf import FPDF

from products.contract.document.contract_parser import parse_contract
from products.contract.domain.document import ParsedDocument


def docx_bytes():
    doc = Document()
    doc.add_heading('第一条 服务内容', 1)
    doc.add_paragraph('保持原文，包含标点。\t制表符\n换行')
    doc.add_table(rows=1, cols=2).cell(0, 0).text = '金额 100'
    doc.add_paragraph('第二条 付款')
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_docx_order_types_exact_text_and_stable_ids():
    version_id = uuid4()
    data = docx_bytes()
    parsed = parse_contract('x.docx', data, version_id=version_id)
    assert isinstance(parsed, ParsedDocument)
    assert parsed.page_count is None
    assert [b.sequence for b in parsed.blocks] == [1, 2, 3, 4]
    assert [b.type for b in parsed.blocks] == ['heading', 'paragraph', 'table', 'paragraph']
    assert all(b.page is None and b.version_id == version_id for b in parsed.blocks)
    assert parsed.blocks[1].text == '保持原文，包含标点。\t制表符\n换行'
    assert all(b.text in parsed.document_text for b in parsed.blocks)
    assert parsed == parse_contract('x.docx', data, version_id=version_id)


def test_pdf_pages_order_and_text():
    pdf = FPDF()
    for label in ['Payment terms', 'Duration one year']:
        pdf.add_page()
        pdf.set_font('Helvetica', size=12)
        pdf.cell(text=label)
    parsed = parse_contract('x.pdf', bytes(pdf.output()))
    assert parsed.page_count == 2
    assert [b.page for b in parsed.blocks] == [1, 2]
    assert [b.sequence for b in parsed.blocks] == [1, 2]
    assert 'Payment terms' in parsed.document_text
    assert parsed.parser_version == 'v2'


def test_empty_and_invalid_documents():
    stream = BytesIO()
    Document().save(stream)
    with pytest.raises(ValueError, match='没有可分析'):
        parse_contract('x.docx', stream.getvalue())
    with pytest.raises(ValueError):
        parse_contract('x.pdf', b'%PDF-invalid')
    with pytest.raises(ValueError):
        parse_contract('x.docx', b'PK-invalid')


def test_merged_table_cell_is_not_duplicated():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = '只出现一次的原文'
    stream = BytesIO()
    doc.save(stream)
    parsed = parse_contract('merged.docx', stream.getvalue())
    assert parsed.blocks[0].text == '只出现一次的原文'


def test_long_paragraph_is_losslessly_split_for_bounded_llm_windows():
    from products.contract.document_intelligence.sections import block_windows
    value = '原文，含有标点。' * 3000
    doc = Document()
    doc.add_paragraph(value)
    stream = BytesIO()
    doc.save(stream)
    parsed = parse_contract('long.docx', stream.getvalue())
    assert len(parsed.blocks) > 1
    assert ''.join(b.text for b in parsed.blocks) == value
    assert len(block_windows(parsed.blocks)) > 1
