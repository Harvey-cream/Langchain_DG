"""Explicit live smoke test with synthetic data; removes only its own fixtures."""
import asyncio
import secrets
import sys
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

import httpx
from unittest.mock import AsyncMock, patch
from docx import Document
from sqlalchemy import delete, select
from app.main import app
from common.database import engine, SessionLocal, create_tables
from common.auth.dependencies import require_user
from common.account.models import User
from products.contract.models import AnalysisRunModel, CustomerModel, ContractModel, ContractVersionModel
from products.contract.document.contract_jobs import execute_analysis
from common.storage.adapter import delete_object
from products.contract import db as contract_db
from products.contract.schemas.document import ClauseCandidate, ClauseExtraction, DocumentMetadata, SectionCandidate, SectionExtraction


class FakeDocumentCapabilities:
    """Mock LLM responses only; source download/parser/store/workflow stay real."""
    async def sections(self, blocks):
        return SectionExtraction(sections=[SectionCandidate(title='正文', source_block_ids=[b.block_id for b in blocks])])

    async def clauses(self, blocks):
        return ClauseExtraction(clauses=[ClauseCandidate(title='合同条款', source_block_ids=[b.block_id for b in blocks])])

    async def metadata(self, blocks):
        return DocumentMetadata(summary='测试合同事实')


async def main():
    await create_tables([User.__table__])
    marker = uuid4().hex
    async with SessionLocal() as db:
        user = User(username='contract-demo-test', email=f'{marker}@example.invalid', password=secrets.token_hex(32))
        db.add(user); await db.commit()
        user_id = user.user_id
    app.dependency_overrides[require_user] = lambda: user
    try:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test', timeout=300) as client:
            customer = (await client.post('/api/customers', json={'name':'演示采购方', 'email':'buyer@example.invalid'})).json()['data']
            contract = (await client.post('/api/contracts', json={'title':'测试采购合同', 'customer_id':customer['id']})).json()['data']
            cid = contract['id']
            doc = Document()
            for paragraph in ['采购合同', '甲方：演示采购方。乙方：演示供应商。', '合同金额为人民币100000元，期限为一年。', '甲方应在签署当日支付全部价款。', '乙方可以随时终止合同且无需退款。', '甲方承担所有间接损失，赔偿金额不设上限。']:
                doc.add_paragraph(paragraph)
            buffer = BytesIO(); doc.save(buffer)
            response = await client.post(f'/api/contracts/{cid}/versions/upload', files={'file':('demo.docx',buffer.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
            assert response.status_code == 200, response.status_code
            vid = response.json()['data']['version_id']
            first_run_id = response.json()['data']['run_id']
            report = (await client.get(f'/api/contracts/{cid}/versions/{vid}/analysis')).json()['data']
            print('Analysis status:', report['status'], flush=True)
            assert report['status'] == 'completed', report.get('error')
            assert report['document_text'] and report['result']['summary']
            async with SessionLocal() as db:
                first_facts = await contract_db.get_document_facts(db, UUID(vid))
                assert first_facts and first_facts.clauses and first_facts.coverage.ratio == 1
                assert all(b.version_id == UUID(vid) for b in first_facts.document.blocks)
            print('Real OSS/DOCX -> Blocks -> Canonical Clauses -> Version Facts passed', flush=True)
            print('Risk count:', len(report['result']['risks']), flush=True)
            repeated = (await client.get(f'/api/contracts/{cid}/versions/{vid}/analysis')).json()['data']
            assert repeated['result'] == report['result']
            listed = (await client.get(f'/api/contracts/{cid}/versions')).json()['data']['versions']
            assert listed[-1]['status'] == 'completed'

            # Two workers may receive the same callback, but only one may claim the run.
            async with SessionLocal() as db:
                claim_run = AnalysisRunModel(version_id=UUID(vid), status='pending')
                db.add(claim_run)
                await db.commit()
                claim_run_id = claim_run.id
            from products.contract.schemas.analysis import ContractAnalysis
            claim_review = AsyncMock(return_value=ContractAnalysis(document_type='测试', summary='原子领取测试'))
            with patch('products.contract.agents.review_agent.analyze_text', claim_review), patch(
                'products.contract.document.contract_source.OssContractDocumentParser.parse',
                AsyncMock(side_effect=AssertionError('cached content should be reused')),
            ):
                await asyncio.gather(execute_analysis(claim_run_id), execute_analysis(claim_run_id))
            assert claim_review.await_count == 1
            async with SessionLocal() as db:
                claimed = await db.get(AnalysisRunModel, claim_run_id)
                claimed_version = await db.get(ContractVersionModel, UUID(vid))
                assert claimed.status == 'completed'
                assert claimed_version.status == 'completed'
            print('Atomic analysis claim and version status passed', flush=True)
            history = (await client.get(f'/api/contracts/{cid}/versions/{vid}/analysis-runs')).json()['data']
            assert len(history['runs']) == 2
            selected = await client.post(
                f'/api/contracts/{cid}/versions/{vid}/analysis-runs/{first_run_id}/select'
            )
            assert selected.status_code == 200
            selected_report = selected.json()['data']
            assert selected_report['selected_run_id'] == first_run_id
            assert selected_report['is_showing_previous'] is True
            print('Analysis history selection passed', flush=True)

            bad_upload = await client.post(f'/api/contracts/{cid}/versions/upload', files={'file': ('bad.pdf', b'invalid', 'application/pdf')})
            assert bad_upload.status_code == 422
            # Exercise concurrent version allocation without extra model charges.
            fake_review = AsyncMock(return_value=ContractAnalysis(document_type='测试', summary='测试分析'))
            with patch('products.contract.agents.review_agent.analyze_text', fake_review):
                uploads = await asyncio.gather(*[
                    client.post(f'/api/contracts/{cid}/versions/upload', files={'file': ('demo.docx', buffer.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
                    for _ in range(2)
                ])
                assert all(r.status_code == 200 for r in uploads)
                listed = (await client.get(f'/api/contracts/{cid}/versions')).json()['data']['versions']
                assert [v['number'] for v in listed] == [1, 2, 3]
                retried = await client.post(f'/api/contracts/{cid}/versions/{vid}/analyze')
                assert retried.status_code == 200
                assert (await client.get(f'/api/contracts/{cid}/versions/{vid}/analysis')).json()['data']['status'] == 'completed'
                async with SessionLocal() as db:
                    assert await contract_db.get_document_facts(db, UUID(vid)) == first_facts
                # Exercise text PDF parsing through the same real upload/job chain.
                from fpdf import FPDF
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font('Helvetica', size=12)
                pdf.cell(text='1. Payment: buyer pays 100 USD within 30 days.')
                uploaded_pdf = await client.post(f'/api/contracts/{cid}/versions/upload',
                    files={'file': ('test.pdf', bytes(pdf.output()), 'application/pdf')})
                assert uploaded_pdf.status_code == 200
                pdf_vid = uploaded_pdf.json()['data']['version_id']
                assert (await client.get(f'/api/contracts/{cid}/versions/{pdf_vid}/analysis')).json()['data']['status'] == 'completed'
                async with SessionLocal() as db:
                    pdf_facts = await contract_db.get_document_facts(db, UUID(pdf_vid))
                    assert pdf_facts.document.page_count == 1
                    assert all(b.page == 1 for b in pdf_facts.document.blocks)
            print('PDF facts persisted; retry preserved Block/Canonical IDs and counts', flush=True)
            print('Invalid upload, concurrent versions, and retry passed', flush=True)
            # A different identity must not retrieve the fixture's text.
            app.dependency_overrides[require_user] = lambda: User(user_id=-1)
            assert (await client.get(f'/api/contracts/{cid}/versions/{vid}/analysis')).status_code == 404
            print('Persistence and ownership checks passed', flush=True)
    finally:
        app.dependency_overrides.pop(require_user, None)
        async with SessionLocal() as db:
            ids = list((await db.scalars(select(ContractModel.id).where(ContractModel.user_id == user_id))).all())
            versions = list((await db.scalars(select(ContractVersionModel).where(ContractVersionModel.contract_id.in_(ids)))).all())
            for version in versions:
                await asyncio.to_thread(delete_object, version.source_key)
            await db.execute(delete(AnalysisRunModel).where(AnalysisRunModel.version_id.in_([v.id for v in versions])))
            await db.execute(delete(ContractVersionModel).where(ContractVersionModel.contract_id.in_(ids)))
            await db.execute(delete(ContractModel).where(ContractModel.user_id == user_id))
            await db.execute(delete(CustomerModel).where(CustomerModel.user_id == user_id))
            await db.execute(delete(User).where(User.user_id == user_id))
            await db.commit()
        await engine.dispose()
        print('Synthetic test records and OSS file removed', flush=True)


if __name__ == '__main__':
    import sys
    if '--mock-review' in sys.argv:
        from products.contract.schemas.analysis import Clause, ContractAnalysis, Risk
        async def fake_review(text):
            return ContractAnalysis(document_type='测试', summary='测试分析',
                clauses=[Clause(title='旧版 Review 条款', original_text=text)],
                risks=[Risk(title='测试风险', risk_level='low', original_text=text,
                    reason='测试原因', suggestion='测试建议', clause_sequence=1)])
        with patch('products.contract.agents.review_agent.analyze_text', AsyncMock(side_effect=fake_review)), patch(
            'products.contract.document.contract_jobs.LangChainDocumentIntelligence', FakeDocumentCapabilities,
        ):
            asyncio.run(main())
    else:
        asyncio.run(main())
